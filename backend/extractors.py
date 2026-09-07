from __future__ import annotations

import re
from pathlib import Path

from .models import DocumentSection

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}
SUPPORTED_MIME_TYPES = {
    ".pdf": {"application/pdf", "application/octet-stream"},
    ".docx": {"application/vnd.openxmlformats-officedocument.wordprocessingml.document", "application/octet-stream"},
    ".txt": {"text/plain", "application/octet-stream"},
    ".md": {"text/markdown", "text/plain", "application/octet-stream"},
    ".markdown": {"text/markdown", "text/plain", "application/octet-stream"},
}


def _markdown_sections(text: str, filename: str) -> list[DocumentSection]:
    sections: list[DocumentSection] = []
    heading = None
    buffer: list[str] = []

    def flush() -> None:
        value = "\n".join(buffer).strip()
        if value:
            sections.append(DocumentSection(source_file=filename, heading=heading, text=value))

    for line in text.splitlines():
        match = re.match(r"^\s{0,3}#{1,6}\s+(.+?)\s*#*\s*$", line)
        if match:
            flush(); buffer.clear(); heading = match.group(1).strip()
        else:
            buffer.append(line)
    flush()
    return sections or [DocumentSection(source_file=filename, text=text.strip())]


def extract_document(path: Path, original_name: str | None = None) -> list[DocumentSection]:
    name = original_name or path.name
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".markdown"}:
        return _markdown_sections(path.read_text(encoding="utf-8", errors="replace"), name)
    if suffix == ".pdf":
        import fitz
        sections: list[DocumentSection] = []
        with fitz.open(path) as document:
            for number, page in enumerate(document, start=1):
                text = page.get_text("text").strip()
                if text:
                    sections.append(DocumentSection(source_file=name, page=number, text=text))
        return sections
    if suffix == ".docx":
        from docx import Document
        document = Document(path)
        sections: list[DocumentSection] = []
        heading = None
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            if paragraph.style and paragraph.style.name.lower().startswith("heading"):
                heading = text
            else:
                sections.append(DocumentSection(source_file=name, heading=heading, text=text))
        for table in document.tables:
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            if rows:
                sections.append(DocumentSection(source_file=name, heading=heading, text="\n".join(rows)))
        return sections
    raise ValueError(f"Unsupported document type: {suffix}")

