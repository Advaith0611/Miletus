from pathlib import Path

from backend.chunking import chunk_sections
from backend.extractors import extract_document
from backend.models import DocumentSection, PodcastScript
from backend.providers import parse_json_object


ROOT = Path(__file__).parents[1]


def test_markdown_extraction_preserves_heading_and_source():
    sections = extract_document(ROOT / "samples" / "chemistry.md", "chemistry.md")
    assert sections[0].heading == "Collision theory"
    assert sections[0].source_file == "chemistry.md"


def test_pdf_extraction_preserves_page_number(tmp_path):
    import fitz
    path = tmp_path / "notes.pdf"
    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 72), "A page of physics notes")
    document.save(path)
    document.close()
    sections = extract_document(path, "notes.pdf")
    assert sections[0].page == 1
    assert "physics" in sections[0].text


def test_docx_extraction_preserves_heading_and_table(tmp_path):
    from docx import Document
    path = tmp_path / "notes.docx"
    document = Document()
    document.add_heading("A chapter", level=1)
    document.add_paragraph("The first idea.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Term"
    table.rows[0].cells[1].text = "Meaning"
    document.save(path)
    sections = extract_document(path, "notes.docx")
    assert sections[0].heading == "A chapter"
    assert any("Term | Meaning" in section.text for section in sections)


def test_chunking_respects_maximum_size_and_metadata():
    sections = [DocumentSection(source_file="notes.txt", page=4, heading="Topic", text="word " * 500)]
    chunks = chunk_sections(sections, max_chars=500)
    assert len(chunks) > 1
    assert all(len(chunk.text) <= 500 for chunk in chunks)
    assert all(chunk.page_start == 4 and chunk.section == "Topic" for chunk in chunks)


def test_script_validation_accepts_two_speakers_and_rejects_missing_speaker():
    valid = {"title": "Test", "segments": [{"id": "1", "speaker": "teacher", "display_text": "Teach."}, {"id": "2", "speaker": "student", "display_text": "Ask."}]}
    assert PodcastScript.model_validate(valid).title == "Test"
    invalid = {"title": "Test", "segments": [{"id": "1", "speaker": "teacher", "display_text": "Teach."}, {"id": "2", "speaker": "teacher", "display_text": "More."}]}
    try:
        PodcastScript.model_validate(invalid)
    except ValueError:
        pass
    else:
        raise AssertionError("missing student speaker should fail")


def test_json_parser_handles_markdown_fences():
    assert parse_json_object('```json\n{"ok": true}\n```') == {"ok": True}
